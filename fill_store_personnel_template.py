#!/usr/bin/env python3
"""
Fill Circle K Store Personnel Safety Walk Excel template with inspection data
"""
import openpyxl
from openpyxl.styles import Font
import sys
import json
import base64
import os

def fill_store_personnel_template(data):
    """Fill the Store Personnel Excel template with safety walk data"""
    
    # Load the template
    template_path = os.path.join(os.path.dirname(__file__), 'Store_Personnel_Safety_Walk_Template.xlsx')
    wb = openpyxl.load_workbook(template_path)
    ws = wb.active
    
    # Fill header info
    ws['B4'] = data.get('name', '')
    ws['E4'] = data.get('storeNumber', '')
    ws['B5'] = data.get('employeeRole', '')
    
    # Combine date and time for display
    date_str = data.get('date', '')
    time_str = data.get('time', '')
    if time_str:
        ws['E5'] = f"{date_str} {time_str} UTC-5"
    else:
        ws['E5'] = date_str
    
    # Questions in order (7 for store personnel)
    questions = [
        "Is the parking lot free of cracks, pot holes, or any other tripping or slipping hazards?",
        "Are all dispensers operational and are hoses and nozzles in good repair?",
        'Are the "Wet Floor" cones/signs visible in all areas during rain and mopping?',
        "Is the dispensed beverage floor area dry and free of any ice and spills?",
        "Are all Dispensed Beverage flavors and CO2 available?",
        "Is the sales floor free of boxes, totes, mop bucket, or anything else a customer could trip over?",
        "The doors are NOT being propped open for a vendor delivery?"
    ]
    
    responses = data.get('responses', {})
    
    # Fill answers starting at row 8
    row = 8
    for question in questions:
        response_data = responses.get(question, {})
        
        # Handle both old format (string) and new format (dict)
        if isinstance(response_data, dict):
            answer = response_data.get('value', '')
            workOrder = response_data.get('workOrder', '')
            correctiveAction = response_data.get('correctiveAction', '')
        else:
            answer = response_data
            workOrder = ''
            correctiveAction = ''
        
        # Fill the cells
        ws[f'C{row}'] = answer
        ws[f'D{row}'] = workOrder
        ws[f'E{row}'] = correctiveAction
        
        # Color code answers
        if answer == 'Yes':
            ws[f'C{row}'].font = Font(bold=True, color='008000', size=11)
        elif answer == 'No':
            ws[f'C{row}'].font = Font(bold=True, color='FF0000', size=11)
        
        row += 1
    
    # Save filled template
    role_short = data.get('employeeRole', '').replace(' ', '')
    filename = f"CircleK_Store{data.get('storeNumber', '')}_{role_short}_{data.get('date', '')}.xlsx"
    filepath = f"/tmp/{filename}"
    wb.save(filepath)
    
    return filepath, filename

if __name__ == '__main__':
    # Read JSON data from stdin
    data_json = sys.stdin.read()
    data = json.loads(data_json)
    
    # Fill the template
    filepath, filename = fill_store_personnel_template(data)
    
    # Read and encode as base64
    with open(filepath, 'rb') as f:
        excel_bytes = f.read()
    
    excel_base64 = base64.b64encode(excel_bytes).decode('utf-8')
    
    # Output result as JSON
    result = {
        'excelBuffer': excel_base64,
        'filename': filename
    }
    print(json.dumps(result))
