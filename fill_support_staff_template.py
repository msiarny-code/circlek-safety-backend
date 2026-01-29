#!/usr/bin/env python3
"""
Fill Circle K Support Staff Safety Walk Excel template with inspection data
"""
import openpyxl
from openpyxl.styles import Font
import sys
import json
import base64
import os

def fill_support_staff_template(data):
    """Fill the Excel template with safety walk data"""
    
    # Load the template
    template_path = '/home/claude/Support_Staff_Safety_Walk_Template.xlsx'
    wb = openpyxl.load_workbook(template_path)
    ws = wb.active
    
    # Fill header info
    ws['B5'] = data.get('name', '')
    ws['E5'] = data.get('storeNumber', '')
    ws['B6'] = data.get('title', '') or data.get('employeeRole', '')
    ws['E6'] = data.get('date', '')
    
    # Questions in order (matching template)
    questions = [
        "Is the parking lot free of cracks, pot holes, or any other tripping or slipping hazards?",
        "Are all dispensers operational and are hoses and nozzles in good repair?",
        'Are the "Wet Floor" cones/signs visible in all areas during rain and mopping? (Please ensure we dry mop often during these times and remove "Wet Floor" cones/signs when area is dry)',
        "Is the dispensed beverage floor area dry and free of any ice and spills?",
        "Are all Dispensed Beverage flavors and CO2 available?",
        "Is the sales floor free of boxes, totes, mop bucket, or anything else a customer could trip over?",
        "Are backroom doors closed or held open with a door stop? (Do not prop open with water or other merchandise)",
        "Is there a clear passage to all exit doors?",
        "Is the area in front of the electrical panels clear of merchandise/boxes?",
        "Are the fire extinguishers hanging and clear of merchandise/boxes?",
        "Is the CO2 detector clear of merchandise/boxes and located near the floor?"
    ]
    
    responses = data.get('responses', {})
    
    # Fill answers starting at row 9
    row = 9
    for question in questions:
        response_data = responses.get(question, {})
        
        # Handle both old format (string) and new format (dict)
        if isinstance(response_data, dict):
            answer = response_data.get('value', '')
            workOrder = response_data.get('workOrder', '')
            correctiveAction = response_data.get('correctiveAction', '')
        else:
            answer = response_data
            workOrder = ''
            correctiveAction = ''
        
        # Fill the cells
        ws[f'C{row}'] = answer
        ws[f'D{row}'] = workOrder
        ws[f'E{row}'] = correctiveAction
        
        # Color code answers
        if answer == 'Yes':
            ws[f'C{row}'].font = Font(bold=True, color='008000', size=11)
        elif answer == 'No':
            ws[f'C{row}'].font = Font(bold=True, color='FF0000', size=11)
        
        row += 1
    
    # Save filled template
    filename = f"CircleK_Store{data.get('storeNumber', '')}_SupportStaff_{data.get('date', '')}.xlsx"
    filepath = f"/tmp/{filename}"
    wb.save(filepath)
    
    return filepath, filename

if __name__ == '__main__':
    # Read JSON data from stdin
    data_json = sys.stdin.read()
    data = json.loads(data_json)
    
    # Fill the template
    filepath, filename = fill_support_staff_template(data)
    
    # Read and encode as base64
    with open(filepath, 'rb') as f:
        excel_bytes = f.read()
    
    excel_base64 = base64.b64encode(excel_bytes).decode('utf-8')
    
    # Output result as JSON
    result = {
        'excelBuffer': excel_base64,
        'filename': filename
    }
    print(json.dumps(result))
